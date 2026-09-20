from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parent / "dsers-family-store-search-terms.docx"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="D9D9D9", size="6"):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    trPr.append(cant_split)

def style_run(run, size=10.5, bold=False, color="000000"):
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def add_para(doc, text="", style=None, space_after=6, size=10.5, bold=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    style_run(r, size=size, bold=bold)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    style_run(p.add_run(text), size=10.5)
    return p

collections = [
    ("Kids and toddler learning", 28, [
        "montessori toddler busy board sensory activity board",
        "toddler fine motor skills toy wooden activity cube",
        "preschool matching puzzle educational toy",
        "toddler pretend play kitchen accessories",
        "kids magnetic tiles building set non electronic",
        "early learning counting sorting toy",
        "toddler travel activity quiet book",
        "kids reusable drawing tablet no battery",]),
    ("Toys and pretend play", 20, [
        "pretend play doctor kit kids non electronic",
        "kids role play tool set",
        "doll stroller pretend play accessory",
        "kids wooden train set",
        "building blocks creative construction set",
        "outdoor bubble toy kids",
        "kids dress up costume accessory",
        "family board game educational",]),
    ("Collectibles and action figures", 16, [
        "collectible action figure display model",
        "posable robot action figure toy",
        "anime style collectible figure generic",
        "manga inspired desk figure original design",
        "blind box style collectible figure",
        "miniature figure display shelf decor",
        "collectible trading card storage binder",
        "model building kit character figure",]),
    ("Anime and manga lifestyle", 14, [
        "anime inspired desk mat original artwork",
        "manga style stationery set original design",
        "anime inspired wall art no trademark",
        "character inspired enamel pin original art",
        "kawaii plush keychain original design",
        "anime inspired tote bag original artwork",
        "manga reading bookmark gift set",
        "otaku room decor original design",]),
    ("Pet essentials for dogs and cats", 22, [
        "slow feeder dog bowl non electric",
        "cat interactive toy no battery",
        "dog grooming slicker brush",
        "pet deshedding comb for dogs cats",
        "washable dog bed cover",
        "cat tunnel foldable indoor",
        "dog training treat pouch",
        "pet travel water bowl non electric",
        "cat scratching mat furniture protector",
        "dog leash collar set reflective non electronic",]),
    ("Baby care and feeding", 18, [
        "silicone baby bib waterproof",
        "baby feeding spoon suction bowl",
        "newborn changing pad portable",
        "baby teething toy food grade silicone",
        "baby bottle drying rack non electric",
        "baby bath rinse cup",
        "infant stroller organizer bag",
        "baby diaper caddy organizer",
        "baby proofing corner guards",]),
    ("Baby wear and kids wear", 24, [
        "baby cotton romper soft newborn",
        "toddler spring fall onesie",
        "kids unisex cotton pajama set",
        "girls summer dress toddler",
        "boys casual outfit set kids",
        "baby muslin swaddle blanket",
        "kids rain jacket lightweight",
        "toddler socks non slip",
        "kids holiday outfit accessory",
        "children sun hat adjustable",]),
    ("Women fashion", 18, [
        "women oversized cotton t shirt",
        "women relaxed fit summer dress",
        "women casual wide leg pants",
        "women lightweight cardigan",
        "women modest blouse everyday",
        "women athleisure sweatshirt",
        "women puffer vest seasonal",
        "women satin skirt outfit",
        "women knit top capsule wardrobe",]),
    ("Women accessories and jewellery", 18, [
        "stainless steel necklace waterproof",
        "minimalist hoop earrings hypoallergenic",
        "layered chain necklace women",
        "fashion brooch lace accessory",
        "women crossbody satchel bag",
        "woven summer tote bag",
        "hair claw clips matte set",
        "silk style hair scarf",
        "women card holder slim wallet",
        "stackable fashion rings stainless steel",]),
    ("Makeup tools and beauty accessories", 16, [
        "makeup brush set synthetic bristles",
        "travel makeup organizer bag",
        "reusable makeup sponge blender",
        "eyelash curler makeup tool",
        "makeup brush cleaning mat",
        "compact cosmetic mirror non electric",
        "lip balm holder keychain",
        "hair styling heatless curlers",]),
    ("Women care skin and hair", 18, [
        "facial cleansing brush manual silicone",
        "skincare headband spa set",
        "satin pillowcase hair care",
        "scalp massage shampoo brush manual",
        "hair detangling brush curly hair",
        "travel toiletry bottles leakproof",
        "under eye cooling mask reusable",
        "body care exfoliating gloves",
        "hair oil applicator comb non electric",]),
    ("Men fashion", 16, [
        "men embroidered short sleeve shirt",
        "men casual overshirt lightweight",
        "men cargo shorts summer",
        "men relaxed fit t shirt",
        "men knit polo shirt",
        "men athletic jogger pants",
        "men puffer vest seasonal",
        "men beach button down shirt",]),
    ("Men accessories", 12, [
        "men leather style card holder",
        "men canvas crossbody bag",
        "men woven belt casual",
        "men minimalist wallet",
        "men sunglasses case",
        "men tie clip lapel pin",
        "men watch strap replacement accessory",
        "men travel toiletry organizer",]),
    ("Home decor and organization", 18, [
        "dopamine decor colorful home accent",
        "afrohemian home decor wall accent",
        "art deco decorative tray",
        "small space storage organizer",
        "wall mounted entryway organizer",
        "decorative candle holder no candle",
        "tabletop vase artificial flower decor",
        "kitchen countertop organizer",
        "bathroom shower storage caddy",]),
    ("Home textiles and handloom", 14, [
        "cotton handloom cushion cover",
        "boho woven throw blanket",
        "reversible quilt bedspread",
        "cotton bedsheet queen size",
        "table runner woven farmhouse",
        "bathroom mat quick dry",
        "decorative pillow cover set",
        "kitchen towel cotton set",]),
    ("Electronics accessories without batteries", 18, [
        "phone case shockproof clear",
        "iphone case magsafe compatible no battery",
        "earbud protective case silicone",
        "wired earbuds microphone",
        "usb c cable braided data cable",
        "laptop stand foldable aluminum",
        "wireless mouse no built in battery excluded power source verify",
        "mechanical keyboard keycap set",
        "webcam privacy cover",]),
    ("Gaming accessories without batteries or logic boards", 10, [
        "wired game controller pc",
        "controller thumb grips silicone",
        "gaming mouse pad extended",
        "keyboard wrist rest ergonomic",
        "gaming headset stand non electronic",
        "console controller carrying case",
        "arcade joystick replacement accessory only if not repair part",
        "game controller wall mount",]),
    ("Watches and watch accessories", 10, [
        "mechanical watch minimalist men",
        "mechanical watch minimalist women",
        "watch organizer travel case",
        "watch display stand",
        "silicone watch strap accessory",
        "stainless steel watch band accessory",
        "watch cleaning cloth kit",
        "watch gift box organizer",]),
    ("Gifts travel and everyday carry", 7, [
        "personalized style gift organizer no personalization claim",
        "travel packing cubes lightweight",
        "weekend travel toiletry bag",
        "reusable shopping tote bag",
        "desk organizer gift set",
        "letter writing stationery gift set",
        "family holiday gift accessory",]),
]

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 12.5)):
    s = styles[name]
    s.font.name = "Aptos Display" if name == "Title" else "Aptos"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(4)
style_run(title.add_run("Family Store DSers Search Terms"), size=22, bold=True)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(14)
style_run(subtitle.add_run("US catalog expansion plan for the remaining Shopify capacity"), size=11, color="555555")

add_para(doc, "Purpose", style="Heading 1", space_after=5, size=16, bold=True)
add_para(doc, "This document gives the DSers search plan and collection allocation for the next catalog expansion pass. The live DSers sidebar currently reads 2,683 products in My Products, so the working capacity to the stated 3,000-product limit is 317 products. The allocation below is a selection target, not a promise to publish every search result: every product must pass the evidence, inventory, safety, and store-fit gates before it is imported or pushed.", size=10.8, space_after=8)
add_para(doc, "The store is positioned as a broad US family retailer: children and baby products, fashion and accessories, beauty tools and care, pets, home, collectibles, and non-battery electronics accessories. Search terms are written to expose useful product intent for classic search and answer-oriented discovery while keeping the final listing grounded in the actual supplier product.", size=10.8, space_after=10)

add_para(doc, "Hard selection rules", style="Heading 1", space_after=5, size=16, bold=True)
for bullet in [
    "Accept only family-store categories and products with a credible US use case.",
    "Require visible supplier stock of at least 200 before selection. Recheck stock immediately before pushing.",
    "Exclude records already marked Pushed to, duplicate supplier listings, and near-identical variants that add no useful assortment.",
    "Exclude batteries, portable chargers, power banks, lithium products, bare logic boards, repair parts, replacement parts, and products whose primary purpose is repair.",
    "Exclude adult or erotic products, unsafe children’s items, medical or therapeutic claims, ingestible products, and cosmetics with unsupported treatment claims.",
    "Reject counterfeit, trademark-infringing, copied character, or unlicensed anime and manga merchandise. Keep anime and manga searches to original or clearly licensed products that DSers evidence supports.",
    "Require clear images, understandable variants, stable cost and shipping information, and a product title that describes what the shopper is actually buying.",
    "For electronics, prefer wired or passive accessories. If a search result includes a battery or a board, skip that result even when the title sounds relevant.",
]:
    add_bullet(doc, bullet)

add_para(doc, "Collection allocation and search terms", style="Heading 1", space_after=5, size=16, bold=True)
add_para(doc, "Use the target column as a portfolio guide while filling the 317 available slots. A collection may receive fewer products when the live results fail the gates; unused capacity should move to another collection with stronger evidence and stock rather than being filled with weak items.", size=10.8, space_after=8)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(1.6), Inches(0.7), Inches(5.0)]
for i, w in enumerate(widths):
    table.columns[i].width = w
headers = ["Collection", "Target", "DSers search terms"]
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    set_cell_shading(cell, "1F4E78")
    set_cell_margins(cell, 120, 120, 120, 120)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    style_run(p.add_run(h), size=10, bold=True, color="FFFFFF")
set_repeat_table_header(table.rows[0])
set_row_cant_split(table.rows[0])
set_table_borders(table)

for idx, (collection, target, terms) in enumerate(collections):
    row = table.add_row()
    set_row_cant_split(row)
    vals = [collection, str(target), "; ".join(terms)]
    for i, val in enumerate(vals):
        cell = row.cells[i]
        cell.width = widths[i]
        set_cell_margins(cell, 100, 120, 100, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx % 2 == 1:
            set_cell_shading(cell, "F3F6F8")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        style_run(p.add_run(val), size=8.6 if i == 2 else 9.3, bold=(i == 0))

add_para(doc, "Search and push sequence", style="Heading 1", space_after=5, size=16, bold=True)
for bullet in [
    "Run one search term at a time in DSers and collect only product cards that show stock at least 200, sensible pricing, and a product-specific image.",
    "Before selecting, open enough of the result to confirm the item is not a battery, power bank, logic board, repair/replacement item, adult product, or unlicensed character product.",
    "De-duplicate by product identity and visual design. Do not fill a collection with many supplier variants that a shopper cannot distinguish.",
    "Add products to the Import List, then push in small verified batches. For each batch, verify the exact Shopify store, push success, and the My Products count increment.",
    "After the batch is pushed, assign the appropriate Shopify collection and run the release workflow. Release verification remains separate from DSers push verification.",
]:
    add_bullet(doc, bullet)

add_para(doc, "Title and description requirements", style="Heading 1", space_after=5, size=16, bold=True)
add_para(doc, "The listing copy must describe the real item in natural language. Put the product type and the shopper’s use case first, then only verified material, size, compatibility, included pieces, and care facts. Do not copy supplier noise such as repeated model numbers, raw catalog tags, shipping origins, or gender labels that the product does not support. Variant names must be readable, such as Black, Blue, 800 by 400 millimeters, or 2 Pack, rather than internal codes.", size=10.8, space_after=8)
add_para(doc, "For answer-oriented discovery, write a short first paragraph that directly answers what the product is, who it is for, and what it helps with. Add a concise specification list and one practical care or compatibility note. Do not make medical, safety, sustainability, or performance claims unless the supplier evidence supports them.", size=10.8, space_after=10)

add_para(doc, "Research basis", style="Heading 1", space_after=5, size=16, bold=True)
add_para(doc, "The category emphasis follows current Shopify merchant trend reporting, including action figures, game controllers, home decor, skin care, satchel bags, body mists, and other fast-moving categories. Home furnishing guidance points to space-saving, colorful decor, outdoor living, smart-home-adjacent accessories, and pet-friendly home products. Google Search Central guidance supports complete merchant product data, clear variant relationships, and consistent product and availability information.", size=10.8, space_after=8)
sources = [
    "Shopify Trending Products: https://www.shopify.com/blog/trending-products",
    "Shopify Home Furnishing Trends: https://www.shopify.com/enterprise/blog/home-furnishing-ecommerce-trends",
    "Google Search Central Product Structured Data: https://developers.google.com/search/docs/appearance/structured-data/product",
    "Shopify Standard Product Taxonomy: https://shopify.github.io/product-taxonomy/releases/2026-02/",
]
for s in sources:
    add_bullet(doc, s)

doc.core_properties.title = "Family Store DSers Search Terms"
doc.core_properties.subject = "US family-store product sourcing and collection allocation"
doc.core_properties.author = "VS Store"
doc.save(OUT)
print(OUT)
